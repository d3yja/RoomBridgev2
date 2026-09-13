from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = Path(__file__).resolve().parents[1]
TITLE = "RoomBridge Preserving Student Needs through Accountable AI Mediation in PolyU Hall Communities"
SUBTHEME = "Sub-theme 3: Social Science, Healthcare Innovations"
SECTIONS = [
    ("Background/Introduction", [
        "Roommate disagreements over sleep, visitors, study routines and shared space make university halls an important setting for examining care in everyday life. An agreement may appear considerate while overlooking a resident's quieter or less strongly expressed requirement. Large language models can produce fluent compromises, but fluency alone cannot establish whether both residents have been heard. When AI supports negotiation, an omitted need can become difficult to notice beneath reassuring language.",
        "RoomBridge addresses this problem within the symposium theme, \u201cCare for life in the AI era\u2014Hall community as a milestone in PolyU's Development 2026.\u201d Its primary alignment is with Social Science, Healthcare Innovations, through attention to resident agency, interpersonal fairness and supportive living environments. The project treats hall community development as a practical opportunity to investigate how AI can assist dialogue while making unresolved concerns visible.",
    ]),
    ("Objective/Research Question", [
        "The study asks whether an explicit process for recording, checking and revising around residents' stated needs reduces their omission from AI-generated roommate agreements compared with simpler mediation approaches. It also examines whether improvements are shared between residents, remain compatible with hall rules and avoid inventing preferences from cultural background. The central hypothesis is that preserving a traceable record of each need, combined with a separate evidence-based audit, will improve coverage and reveal compromises requiring further discussion.",
    ]),
    ("Methods", [
        "The evaluation will use the prototype's 22 authored synthetic scenarios, covering everyday roommate tensions and situations requiring escalation. Each scenario contains a fixed reference set of residents' stated needs, including explicit boundaries and importance where provided. Four conditions will be compared: generic AI mediation, context-informed mediation, multi-agent deliberation and the complete RoomBridge workflow. Repeated runs with specified random seeds and recorded model identities will support reproducibility. Comparisons will be paired within scenarios, with uncertainty reported across scenarios and repeated runs.",
        "RoomBridge represents each need separately, identifies conflicts and generates candidate agreements. A separate auditor checks each need against the agreement's wording without seeing the mediator's rationale or deliberation. Positive coverage judgements require supporting quotations from the agreement. Missing, contradicted or unresolved needs trigger a bounded revision process. Hall-rule checks separately identify prohibited terms and tensions between stated needs and institutional requirements. Contextual explanations remain attached to stated needs; assumption checks flag preferences or facts that residents never supplied.",
        "Primary outcomes will include Needs-Retention Rate, a graded measure of how well agreements address stated needs; Silent Loss Rate, the proportion omitted or contradicted; and Retention Asymmetry, the difference between residents' coverage scores. Policy violations and escalation behaviour will provide additional checks. Because conditions differ in access to hall rules, a matched-information comparison is proposed to examine whether gains reflect the workflow or additional information. Independent human reviewers will label a held-out sample of agreements to assess automated judgements and examine disputed cases.",
        "Preservation does not mean promising that every preference can be satisfied. For example, a request to host visitors may conflict with another resident's sleep schedule or a binding hall rule. An agreement should identify that conflict, explain what remains unresolved and invite residents to decide acceptable alternatives. This matters because a high coverage score alone cannot show that a settlement is feasible, voluntarily accepted or fair. The evaluation will therefore examine unresolved cases alongside numerical scores, rather than treating completion as successful mediation.",
        "Threats, coercion and other serious concerns will be evaluated as escalation cases, where agreement generation should stop. Any later evaluation involving residents would require appropriate ethics review, informed consent and protection of personal information. Residents would retain the ability to reject suggestions and seek human support.",
    ]),
    ("Results/Findings", [
        "The working prototype implements the four conditions, structured needs records, agreement auditing, revision and escalation checks. Deterministic offline tests demonstrate that deliberately omitted needs can be detected and recovered in the designed test cases, and that threat scenarios can halt mediation. These findings establish functional feasibility within a controlled demonstration. They do not establish comparative effectiveness of live language models or improvements in resident wellbeing. The planned model evaluation will test whether these mechanisms remain useful under variable generated responses, while human review will examine whether the recorded coverage corresponds to meaningful acknowledgement of residents' concerns.",
    ]),
    ("Conclusion/Implications", [
        "RoomBridge proposes a concrete standard for caring AI in hall life: every stated need should remain inspectable, and unresolved tensions should be surfaced for discussion. Its contribution is a testable connection between accountable AI design and interpersonal fairness. If evaluation supports the approach, hall staff could use its records to support more balanced conversations and identify previously overlooked concerns. For PolyU's development in 2026, this offers a focused route towards sustainable hall communities where technological assistance strengthens resident participation and preserves human responsibility for difficult decisions.",
    ]),
]
KEYWORDS = "AI mediation; student halls; needs preservation; interpersonal fairness; responsible AI"

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.2677)
section.page_height = Inches(11.6929)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.9)
normal = doc.styles['Normal']
normal.font.name = 'Times New Roman'
normal.font.size = Pt(12)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.line_spacing = 1.1
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.widow_control = True
for name in ['Title', 'Subtitle', 'Heading 1', 'Heading 2']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(12)
    style.paragraph_format.space_before = Pt(7)
    style.paragraph_format.space_after = Pt(3)
    if name.startswith('Heading'):
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
doc.styles['Title'].font.size = Pt(15)
doc.styles['Title'].font.bold = True
doc.styles['Title'].paragraph_format.space_after = Pt(6)
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    if style.name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2']:
        fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
        for key in list(fonts.attrib):
            if key.endswith('Theme'):
                del fonts.attrib[key]
        fonts.set(qn('w:ascii'), 'Times New Roman')
        fonts.set(qn('w:hAnsi'), 'Times New Roman')
doc.add_paragraph(TITLE, 'Title')
p = doc.add_paragraph(SUBTHEME)
p.paragraph_format.space_after = Pt(8)
p.runs[0].italic = True
doc.add_paragraph('Abstract', 'Heading 1')
for heading, paragraphs in SECTIONS:
    doc.add_paragraph(heading, 'Heading 2')
    for text in paragraphs:
        p = doc.add_paragraph(text)
        if text.startswith('Primary outcomes'):
            p.paragraph_format.page_break_before = True
p = doc.add_paragraph()
p.add_run('Keywords: ').bold = True
p.add_run(KEYWORDS)
for run in p.runs:
    run.font.size = Pt(11)
doc.core_properties.title = TITLE
doc.core_properties.subject = 'PolyU hall symposium research proposal abstract'
doc.core_properties.author = ''
doc.core_properties.keywords = KEYWORDS
text = '\n'.join(p.text for p in doc.paragraphs)
print('Total words:', len(text.split()))
print('Title words:', len(TITLE.split()))
for heading, paragraphs in SECTIONS:
    print(heading, sum(len(t.split()) for t in paragraphs))
(OUT / 'qa' / 'abstract_text.txt').write_text(text, encoding='utf-8')
doc.save(OUT / 'RoomBridge_Symposium_Abstract.docx')
