from pathlib import Path
import sys
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parents[1]
TITLE = 'RoomBridge Supporting Mental Health and Wellbeing through Accountable AI Mediation in PolyU Halls'
KEYWORDS = 'AI mediation; student mental health; wellbeing; roommate relationships; needs preservation'
SECTIONS = [
    ('a. Background/Introduction', [
        "Shared hall life places sleep, privacy and interpersonal relationships at the centre of students' everyday wellbeing. Foulkes et al. (2021) found that relationships within university accommodation could provide support or contribute to loneliness and distress. Stores et al. (2023) reported associations between sleep disturbance, mental health and wellbeing among university students. These findings motivate examining roommate negotiation as one possible way to support living conditions relevant to mental health.",
        "AI-generated agreements may sound considerate while omitting a resident's stated requirement, such as uninterrupted sleep or time alone. RoomBridge investigates this potential failure by keeping each need visible throughout mediation. Aligned with Sub-theme 3, Social Science, Healthcare Innovations, the proposal connects the symposium's care for life in the AI era theme to respectful dialogue and supportive hall environments at PolyU.",
    ]),
    ('b. Objectives', [
        "The primary question is whether recording, auditing and revising around stated needs reduces omissions from AI-mediated roommate agreements compared with simpler approaches. A secondary question asks whether residents experience greater acknowledgement, fairness and control over proposed arrangements. The hypothesised pathway is that clearer agreements and visible unresolved concerns could reduce conflict-related stress and protect sleep. Improved need coverage will be evaluated separately from mental health benefits, which require evidence from residents.",
        "The proposed innovation is to make care assessable at the level of individual needs. For example, an agreement about visitors should explicitly address a roommate's sleep requirement, record any remaining disagreement and allow both residents to reconsider the arrangement. A polished compromise would therefore be insufficient evidence of success. Resident feedback would help distinguish agreements that merely mention a need from arrangements that feel workable, respectful and supportive in daily shared hall life at PolyU.",
    ]),
    ('c. Methods', [
        "The initial evaluation will compare four conditions across 22 authored synthetic scenarios: generic AI mediation, context-informed mediation, multi-agent deliberation and the complete RoomBridge workflow. Each scenario has a fixed reference set of stated needs. Repeated runs will record models, prompts and random seeds, with comparisons paired within scenarios. A matched-information comparison is proposed to account for differences in access to hall rules.",
        "RoomBridge records needs individually, identifies conflicts and generates candidate agreements. A separate auditor examines each need against the agreement without seeing the mediator's reasoning. Positive coverage judgements require quotations from the agreement. Missing, contradicted or unresolved needs trigger bounded revision. Hall-rule checks identify prohibited arrangements, while assumption checks flag invented preferences or cultural stereotypes. Conflicting needs remain visible even when satisfying both is impossible.",
        "Primary measures will assess graded need coverage, the proportion of needs omitted or contradicted, and differences in coverage between residents. Independent human reviewers will assess a held-out sample to check automated judgements. Threats and serious safety concerns will be tested as escalation cases requiring mediation to stop.",
        "A subsequent exploratory pilot with consenting residents is proposed, subject to ethics approval. Questionnaires before mediation and at follow-up would assess perceived fairness, feeling heard, conflict-related stress, sleep disruption and wellbeing, using validated measures where appropriate. Interviews would examine acceptability and whether agreements were voluntarily adopted. A comparison group receiving usual hall support would help interpret changes. Participation would remain voluntary, personal information would be minimised, and residents could reject suggestions or seek qualified human support.",
    ]),
    ('d. Expected Results/Findings', [
        "The prototype already implements structured needs records, auditing, revision and escalation. Offline tests show that deliberately omitted needs can be recovered in designed cases and that threat scenarios can halt mediation. These are functional demonstrations. The planned evaluation will test whether RoomBridge improves coverage and reduces unequal omissions with live language models. The resident pilot would explore whether any improvements accompany lower stress, less sleep disruption or better wellbeing. No clinical effectiveness or mental health improvement has yet been established.",
    ]),
    ('e. Conclusion/Discussion, Implications, and Future Directions', [
        "RoomBridge frames care as keeping residents' needs visible and preserving their authority over shared living decisions. Its proposed contribution is an accountable mediation process with measurable links to everyday wellbeing. If supported by evaluation, hall staff could use its records to identify overlooked concerns and guide timely support. The system is intended to assist communication, with counselling and clinical care remaining the responsibility of qualified professionals. This approach positions hall community development in PolyU's 2026 vision as an opportunity to investigate responsible AI that supports healthier, more sustainable relationships.",
    ]),
]
REFERENCES = [
    ('Foulkes, L., Reddy, A., Westbrook, J., Newbronner, E., & McMillan, D. (2021). Social relationships within university undergraduate accommodation: A qualitative study. Journal of Further and Higher Education, 45(10), 1469–1482.', 'https://doi.org/10.1080/0309877X.2021.1879745'),
    ('Stores, R., Linceviciute, S., Pilkington, K., & Ridge, D. (2023). Sleep disturbance, mental health, wellbeing and educational impact in UK university students: A mixed methods study. Journal of Further and Higher Education, 47(8), 995–1008.', 'https://doi.org/10.1080/0309877X.2023.2209777'),
]

stem = 'RoomBridge_Wellbeing_Abstract'
referenced = len(sys.argv) > 1 and sys.argv[1] == 'referenced'
if referenced:
    from referenced_content import TITLE, KEYWORDS, SECTIONS, REFERENCES
    stem = 'RoomBridge_Referenced_Abstract'

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.2677)
section.page_height = Inches(11.6929)
section.top_margin = section.bottom_margin = Inches(0.8)
section.left_margin = section.right_margin = Inches(0.9)
for name in ['Normal', 'Title', 'Heading 1']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(11.5 if referenced else 12)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for key in list(fonts.attrib):
        if key.endswith('Theme'):
            del fonts.attrib[key]
    fonts.set(qn('w:ascii'), 'Times New Roman')
    fonts.set(qn('w:hAnsi'), 'Times New Roman')
for style in doc.styles:
    for border in list(style.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.08
normal.space_after = Pt(6)
normal.widow_control = True
heading = doc.styles['Heading 1']
heading.font.bold = True
heading.paragraph_format.space_before = Pt(7)
heading.paragraph_format.space_after = Pt(2)
heading.paragraph_format.keep_with_next = True
title = doc.styles['Title']
title.font.size = Pt(14)
title.font.bold = True
title.paragraph_format.space_after = Pt(8)
doc.add_paragraph(TITLE, 'Title')
p = doc.add_paragraph()
p.add_run('Keywords: ').bold = True
p.add_run(KEYWORDS)
for heading, paragraphs in SECTIONS:
    h = doc.add_paragraph(heading, 'Heading 1')
    if referenced and heading.startswith('d.'):
        h.paragraph_format.page_break_before = True
    for text in paragraphs:
        p = doc.add_paragraph(text)
        if not referenced and text.startswith('Primary measures'):
            p.paragraph_format.page_break_before = True
doc.add_paragraph('References', 'Heading 1')
for reference, url in REFERENCES:
    p = doc.add_paragraph(reference + ' ')
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.first_line_indent = Inches(-0.18)
    p.paragraph_format.line_spacing = 1.0
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    for run in p.runs:
        run.font.size = Pt(10)
    link = OxmlElement('w:hyperlink')
    link.set(qn('r:id'), p.part.relate_to(url, RT.HYPERLINK, is_external=True))
    run = OxmlElement('w:r')
    props = OxmlElement('w:rPr')
    size = OxmlElement('w:sz')
    size.set(qn('w:val'), '20')
    props.append(size)
    run.append(props)
    text = OxmlElement('w:t')
    text.text = url
    run.append(text)
    link.append(run)
    p._p.append(link)
doc.core_properties.title = TITLE
doc.core_properties.subject = 'PolyU symposium proposal on AI mediation and student wellbeing'
doc.core_properties.author = ''
doc.core_properties.keywords = KEYWORDS
all_text = '\n'.join(''.join(t.text or '' for t in p._p.iter(qn('w:t'))) for p in doc.paragraphs)
print('Total words:', len(all_text.split()))
print('Title words:', len(TITLE.split()))
print('Body words:', sum(len(t.split()) for _, ps in SECTIONS for t in ps))
(OUT / 'qa' / f'{stem}_text.txt').write_text(all_text, encoding='utf-8')
doc.save(OUT / f'{stem}.docx')
